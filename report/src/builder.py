"""Render the thesis content (a list of blocks) to a .docx that follows the IIT Kharagpur report guideline.

Block types (tuples):
  ("chapter", key, title)            new chapter, numbered 1, 2, ... (or "Appendix A" when appendix mode is on)
  ("appendix_mode",)                 following chapters are appendices A, B, ...
  ("h2", key, title) / ("h3", key, title)
  ("p", text)                        body paragraph with inline markup (see parse_inline)
  ("bullets", [text, ...]) / ("numbered", [text, ...])
  ("eq", key, latex)                 display equation, numbered (n.k)
  ("fig", key, path, width_in, caption)
  ("table", key, caption, header, rows, col_widths_in)
  ("pagebreak",)
Inline markup: **bold**, *italic*, ^{sup}, _{sub}, [@k1,k2] citations, {fig:k} {tab:k} {eq:k} {sec:k} {ch:k} refs.
"""
import re
import io
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from PIL import Image
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

FONT = "Times New Roman"
TEXT_WIDTH_IN = 6.0  # 8.5 in - 1.5 in binding - 1.0 in outer
TAB_RIGHT_IN = 5.95  # right-aligned tab stops sit just inside the margin, or Pages wraps them onto a new line
MINUS = "\u2212"


# ------------------------------------------------------------------ numbering pre-pass
class Numbering:
    def __init__(self, blocks, refs_order_keys):
        self.labels, self.cite_order = {}, []
        appendix, ch, letters = False, 0, iter("ABCDEFGH")
        prefix = None
        h2 = h3 = fig = tab = eq = 0
        for b in blocks:
            kind = b[0]
            if kind == "appendix_mode":
                appendix = True
                continue
            if kind == "chapter":
                if appendix:
                    prefix = next(letters)
                else:
                    ch += 1
                    prefix = str(ch)
                h2 = h3 = fig = tab = eq = 0
                self.labels[b[1]] = prefix
            elif kind == "h2":
                h2 += 1; h3 = 0
                self.labels[b[1]] = f"{prefix}.{h2}"
            elif kind == "h3":
                h3 += 1
                self.labels[b[1]] = f"{prefix}.{h2}.{h3}"
            elif kind == "fig":
                fig += 1; self.labels[b[1]] = f"{prefix}.{fig}"
            elif kind == "table":
                tab += 1; self.labels[b[1]] = f"{prefix}.{tab}"
            elif kind == "eq":
                eq += 1; self.labels[b[1]] = f"{prefix}.{eq}"
            for text in block_texts(b):
                for group in re.findall(r"\[@([^\]]+)\]", text):
                    for k in [x.strip() for x in group.split(",")]:
                        if k not in self.cite_order:
                            self.cite_order.append(k)
        unknown = [k for k in self.cite_order if k not in refs_order_keys]
        if unknown:
            raise KeyError(f"citations without a reference entry: {unknown}")
        self.cite_num = {k: i + 1 for i, k in enumerate(self.cite_order)}


def block_texts(b):
    kind = b[0]
    if kind == "p":
        return [b[1]]
    if kind in ("bullets", "numbered"):
        return list(b[1])
    if kind == "fig":
        return [b[4]]
    if kind == "table":
        return [b[2]] + [str(c) for row in b[4] for c in row]
    return []


# ------------------------------------------------------------------ low-level XML helpers
def insert_before_any(parent, el, successors):
    """Insert el before the first existing child named in successors (schema order), else append."""
    for name in successors:
        found = parent.find(qn(name))
        if found is not None:
            found.addprevious(el)
            return el
    parent.append(el)
    return el


TBLPR_AFTER_BORDERS = ["w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook", "w:tblCaption", "w:tblDescription"]
def looks_numeric(v):
    """True for cells that read as numbers (centred in tables): 0.123, +0.16, < 0.001, 72%, 0.86 / -0.33 / 67%, best."""
    v = v.strip().replace("**", "")
    return bool(v) and bool(re.fullmatch(r"(best|[<>≤≥]?\s*[+\-−]?\d[\d.,]*\s*%?(\s*[(/]\s*[+\-−]?\d[\d.,]*\s*%?\)?)*)", v))


def set_booktabs(table, n_rows, heavy=12, light=6):
    """Academic three-rule table: heavy rule on top and bottom, light rule under the header, nothing else."""
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        if edge in ("top", "bottom"):
            el.set(qn("w:val"), "single"); el.set(qn("w:sz"), str(heavy)); el.set(qn("w:color"), "000000")
        else:
            el.set(qn("w:val"), "nil")
        borders.append(el)
    tblPr.append(borders)
    for cell in table.rows[0].cells:
        tcPr = cell._tc.get_or_add_tcPr()
        tcb = OxmlElement("w:tcBorders"); b = OxmlElement("w:bottom")
        b.set(qn("w:val"), "single"); b.set(qn("w:sz"), str(light)); b.set(qn("w:color"), "000000")
        tcb.append(b); tcPr.append(tcb)


def set_grid_widths(table, widths_in):
    """Write the column widths into w:tblGrid as well; several viewers ignore per-cell widths."""
    grid = table._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths_in):
        gc.set(qn("w:w"), str(int(w * 1440)))
    tblPr = table._tbl.tblPr
    tw = tblPr.find(qn("w:tblW"))
    if tw is None:
        tw = OxmlElement("w:tblW"); tblPr.append(tw)
    tw.set(qn("w:w"), str(int(sum(widths_in) * 1440))); tw.set(qn("w:type"), "dxa")
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)


def set_cell_valign(cell, val="center"):
    tcPr = cell._tc.get_or_add_tcPr()
    v = OxmlElement("w:vAlign"); v.set(qn("w:val"), val); tcPr.append(v)


def set_cell_shading(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def set_cell_margins(table, top=40, bottom=40, left=80, right=80):
    tblPr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for side, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{side}"); el.set(qn("w:w"), str(v)); el.set(qn("w:type"), "dxa"); mar.append(el)
    insert_before_any(tblPr, mar, ["w:tblLook", "w:tblCaption", "w:tblDescription"])


def set_table_borders(table, sz=6):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}"); el.set(qn("w:val"), "single"); el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), "0"); el.set(qn("w:color"), "808080"); borders.append(el)
    insert_before_any(tblPr, borders, TBLPR_AFTER_BORDERS)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader"); el.set(qn("w:val"), "true"); trPr.append(el)


def no_split(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit"); el.set(qn("w:val"), "true"); trPr.append(el)


def add_field(paragraph, instr, placeholder=""):
    """Simple field such as PAGE inside a paragraph."""
    r = paragraph.add_run(); fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "begin"); r._r.append(fc)
    r = paragraph.add_run(); it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr; r._r.append(it)
    r = paragraph.add_run(); fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "separate"); r._r.append(fc)
    paragraph.add_run(placeholder)
    r = paragraph.add_run(); fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "end"); r._r.append(fc)


def field_begin(paragraph, instr):
    r = paragraph.add_run(); fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "begin"); fc.set(qn("w:dirty"), "true"); r._r.append(fc)
    r = paragraph.add_run(); it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr; r._r.append(it)
    r = paragraph.add_run(); fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "separate"); r._r.append(fc)


def field_end(paragraph):
    r = paragraph.add_run(); fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "end"); r._r.append(fc)


def set_page_numbering(section, fmt=None, start=None):
    sectPr = section._sectPr
    pg = sectPr.find(qn("w:pgNumType"))
    if pg is None:
        pg = OxmlElement("w:pgNumType")
        insert_before_any(sectPr, pg, ["w:cols", "w:formProt", "w:vAlign", "w:noEndnote", "w:titlePg",
                                       "w:textDirection", "w:bidi", "w:rtlGutter", "w:docGrid", "w:printerSettings"])
    if fmt:
        pg.set(qn("w:fmt"), fmt)
    if start is not None:
        pg.set(qn("w:start"), str(start))


def add_bookmark(paragraph, name, bid):
    start = OxmlElement("w:bookmarkStart"); start.set(qn("w:id"), str(bid)); start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd"); end.set(qn("w:id"), str(bid))
    pPr = paragraph._p.find(qn("w:pPr"))
    if pPr is not None:
        pPr.addnext(start)
    else:
        paragraph._p.insert(0, start)
    paragraph._p.append(end)


# ------------------------------------------------------------------ equations as images
def render_equation(latex, out_png, fontsize=13):
    fig = plt.figure(figsize=(0.01, 0.01))
    fig.text(0, 0, f"${latex}$", fontsize=fontsize)
    plt.rcParams["mathtext.fontset"] = "stix"
    fig.savefig(out_png, dpi=300, bbox_inches="tight", pad_inches=0.03, transparent=False, facecolor="white")
    plt.close(fig)
    with Image.open(out_png) as im:
        return im.size[0] / 300, im.size[1] / 300


# ------------------------------------------------------------------ the builder
class ThesisBuilder:
    def __init__(self, blocks, refs, eq_dir, pages=None):
        self.blocks, self.refs = blocks, refs
        self.num = Numbering(blocks, set(refs))
        self.eq_dir = Path(eq_dir); self.eq_dir.mkdir(parents=True, exist_ok=True)
        self.pages = pages or {}  # label text -> printed page number, filled on the second pass
        self.doc = Document()
        self.toc_entries, self.lof_entries, self.lot_entries = [], [], []
        self._bm = 100
        self._setup_styles()

    # ---------------- styles and page setup
    def _setup_styles(self):
        d = self.doc
        sec = d.sections[0]
        sec.page_width, sec.page_height = Inches(8.5), Inches(11)
        sec.left_margin, sec.right_margin = Inches(1.5), Inches(1.0)
        sec.top_margin, sec.bottom_margin = Inches(1.0), Inches(1.0)
        sec.footer_distance = Inches(0.5); sec.header_distance = Inches(0.5)
        st = d.styles
        normal = st["Normal"]
        normal.font.name = FONT; normal.font.size = Pt(12)
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        pf = normal.paragraph_format
        pf.line_spacing = 1.5; pf.space_after = Pt(10); pf.space_before = Pt(0)
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY; pf.widow_control = True
        for name, size, bold, italic, before, after in [("Heading 1", 18, True, False, 24, 30), ("Heading 2", 13, True, False, 16, 6),
                                                        ("Heading 3", 12, True, True, 12, 4)]:
            s = st[name]
            s.font.name = FONT; s.font.size = Pt(size); s.font.bold = bold; s.font.italic = italic
            s.font.color.rgb = RGBColor(0, 0, 0)
            rpr = s.element.get_or_add_rPr()
            rfonts = rpr.find(qn("w:rFonts"))
            if rfonts is None:
                rfonts = OxmlElement("w:rFonts"); rpr.append(rfonts)
            for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
                rfonts.set(qn(a), FONT)
            s.paragraph_format.space_before = Pt(before); s.paragraph_format.space_after = Pt(after)
            s.paragraph_format.line_spacing = 1.15; s.paragraph_format.keep_with_next = True
            s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        st["Heading 1"].paragraph_format.page_break_before = True
        st["Heading 1"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        st["Heading 1"].paragraph_format.line_spacing = 1.3
        for name in ("Heading 2", "Heading 3"):
            hp = st[name].paragraph_format
            hp.left_indent = Inches(0.6); hp.first_line_indent = Inches(-0.6)
            hp.tab_stops.add_tab_stop(Inches(0.6))
        # built-in TOC styles, used when Word or Pages regenerates the contents: right-aligned page numbers with dots
        for lvl, ind in ((1, 0.0), (2, 0.35), (3, 0.75)):
            ts = st.add_style(f"TOC {lvl}", 1) if f"TOC {lvl}" not in [x.name for x in st] else st[f"TOC {lvl}"]
            ts.base_style = normal; ts.font.size = Pt(12); ts.font.bold = False
            tp = ts.paragraph_format
            tp.line_spacing = 1.0; tp.space_after = Pt(4); tp.space_before = Pt(8 if lvl == 1 else 0)
            tp.left_indent = Inches(ind); tp.right_indent = Inches(0); tp.alignment = WD_ALIGN_PARAGRAPH.LEFT
            tp.tab_stops.add_tab_stop(Inches(TAB_RIGHT_IN), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        for cname in ("Figure Caption", "Table Caption"):
            s = st.add_style(cname, 1)
            s.base_style = normal; s.font.size = Pt(11)
            s.paragraph_format.line_spacing = 1.0; s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            s.paragraph_format.space_after = Pt(14 if cname == "Figure Caption" else 6)
            s.paragraph_format.space_before = Pt(4 if cname == "Figure Caption" else 10)
        st["Table Caption"].paragraph_format.keep_with_next = True
        s = st.add_style("Table Text", 1); s.base_style = normal; s.font.size = Pt(10)
        s.paragraph_format.line_spacing = 1.0; s.paragraph_format.space_after = Pt(0); s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        s = st.add_style("Reference", 1); s.base_style = normal; s.font.size = Pt(11)
        s.paragraph_format.line_spacing = 1.0; s.paragraph_format.space_after = Pt(6)
        s.paragraph_format.left_indent = Inches(0.45); s.paragraph_format.first_line_indent = Inches(-0.45)
        s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        s = st.add_style("Single", 1); s.base_style = normal; s.paragraph_format.line_spacing = 1.0
        s = st.add_style("TOC Entry", 1); s.base_style = normal; s.paragraph_format.line_spacing = 1.0
        s.paragraph_format.space_after = Pt(5); s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        s.paragraph_format.right_indent = Inches(0)
        for lst in ("List Bullet", "List Number"):
            ls = st[lst]; ls.font.name = FONT; ls.font.size = Pt(12)
            ls.paragraph_format.line_spacing = 1.5; ls.paragraph_format.space_after = Pt(4)
            ls.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        # ask Word to refresh the Contents / List of Figures / List of Tables fields when the file is opened
        settings = d.settings.element
        zoom = settings.find(qn("w:zoom"))
        if zoom is not None and zoom.get(qn("w:percent")) is None:
            zoom.set(qn("w:percent"), "100")
        uf = OxmlElement("w:updateFields"); uf.set(qn("w:val"), "true")
        insert_before_any(settings, uf, ["w:hdrShapeDefaults", "w:footnotePr", "w:endnotePr", "w:compat", "w:docVars",
                                         "w:rsids", "w:attachedSchema", "w:themeFontLang", "w:clrSchemeMapping",
                                         "w:doNotIncludeSubdocsInStats", "w:doNotAutoCompressPictures", "w:forceUpgrade",
                                         "w:captions", "w:readModeInkLockDown", "w:smartTagType",
                                         "w:shapeDefaults", "w:doNotEmbedSmartTags", "w:decimalSymbol", "w:listSeparator"])

    # ---------------- inline text
    def resolve(self, text):
        def cite(m):
            nums = sorted(self.num.cite_num[k.strip()] for k in m.group(1).split(","))
            # compress runs: [3]-[5]
            out, i = [], 0
            while i < len(nums):
                j = i
                while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1:
                    j += 1
                out.append(f"[{nums[i]}]–[{nums[j]}]" if j - i >= 2 else ", ".join(f"[{n}]" for n in nums[i:j + 1]))
                i = j + 1
            return ", ".join(out)
        text = re.sub(r"\[@([^\]]+)\]", cite, text)
        def ref(m):
            kind, key = m.group(1), m.group(2)
            lab = self.num.labels[f"{kind}:{key}"]
            if kind == "ch" and lab.isalpha():
                return "Appendix " + lab
            return {"fig": "Figure ", "tab": "Table ", "eq": "Equation (", "sec": "Section ", "ch": "Chapter "}[kind] + lab + (")" if kind == "eq" else "")
        text = re.sub(r"\(Section \{sec:", "({sec:", text)
        text = re.sub(r"\bSection \{sec:", "{sec:", text)
        text = re.sub(r"\{(fig|tab|eq|sec|ch):([\w\-]+)\}", ref, text)
        text = re.sub(r"\b(Table|Figure) ([\w.]+)–(?:Table|Figure) ([\w.]+)", r"\1s \2–\3", text)
        # typographic minus for negative numbers (not hyphenated words, ISO dates or command-line flags)
        text = re.sub(r"(?<![\w\-])-(?=\d)", MINUS, text)
        # keep numbers with their units on one line (no stray "h" or "m" at the start of a line)
        return re.sub(r"(\d) (?=(h|m|km|hPa|m/s|s|°C|K|%)\b)", "\\1\u00a0", text)

    def add_inline(self, paragraph, text, size=None, bold=False, italic=False):
        text = self.resolve(text)
        tokens = re.split(r"(\*\*.+?\*\*|\*.+?\*|\^\{.+?\}|_\{.+?\}|`.+?`)", text)
        for tok in tokens:
            if not tok:
                continue
            b, i, sup, sub, mono = bold, italic, False, False, False
            if tok.startswith("`") and tok.endswith("`") and len(tok) > 2:
                tok, mono = tok[1:-1], True
            elif tok.startswith("**") and tok.endswith("**"):
                tok, b = tok[2:-2], True
            elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
                tok, i = tok[1:-1], True
            elif tok.startswith("^{"):
                tok, sup = tok[2:-1], True
            elif tok.startswith("_{"):
                tok, sub = tok[2:-1], True
            r = paragraph.add_run(tok)
            r.bold, r.italic = b or None, i or None
            if sup: r.font.superscript = True
            if sub: r.font.subscript = True
            if size: r.font.size = Pt(size)
            r.font.name = FONT
            if mono:
                r.font.name = "Courier New"; r.font.size = Pt((size or 12) - 1)
        return paragraph

    def para(self, text, style=None, align=None, size=None, bold=False, italic=False, space_after=None, keep_next=False):
        p = self.doc.add_paragraph(style=style)
        self.add_inline(p, text, size=size, bold=bold, italic=italic)
        if align is not None: p.paragraph_format.alignment = align
        if space_after is not None: p.paragraph_format.space_after = Pt(space_after)
        if keep_next: p.paragraph_format.keep_with_next = True
        return p

    def page_of(self, key):
        return self.pages.get(key, "0")

    # ---------------- block renderers
    def heading(self, level, text, key):
        p = self.doc.add_paragraph(style=f"Heading {level}")
        if level == 1 and "  " in text:
            # "Chapter 3  Study Area and Data" -> small spaced label above the title, both centred
            label, title = text.split("  ", 1)
            r = p.add_run(label.upper()); r.font.size = Pt(13); r.bold = True
            r.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
            r.add_break()
            self.add_inline(p, title)
        elif level in (2, 3) and "  " in text:
            self.add_inline(p, text.replace("  ", "\t", 1))
        else:
            self.add_inline(p, text.upper() if level == 1 else text)
        self._bm += 1
        add_bookmark(p, f"_Toc{self._bm}", self._bm)
        self.toc_entries.append((level, text, key))
        return p

    def figure(self, key, path, width, caption):
        p = self.doc.add_paragraph(style="Single")
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(2)
        p.add_run().add_picture(str(path), width=Inches(width))
        lab = self.num.labels[key]
        short, caption = caption.split("|", 1) if "|" in caption else (None, caption)
        c = self.doc.add_paragraph(style="Figure Caption")
        r = c.add_run(f"Figure {lab}: "); r.bold = True
        self.add_inline(c, caption)
        self.lof_entries.append((f"Figure {lab}", short or caption.split(". ")[0].rstrip("."), key))

    def table(self, key, caption, header, rows, widths):
        lab = self.num.labels[key]
        short, caption = caption.split("|", 1) if "|" in caption else (None, caption)
        c = self.doc.add_paragraph(style="Table Caption")
        r = c.add_run(f"Table {lab}: "); r.bold = True
        self.add_inline(c, caption)
        self.lot_entries.append((f"Table {lab}", short or caption.split(". ")[0].rstrip("."), key))
        t = self.doc.add_table(rows=1 + len(rows), cols=len(header))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
        set_booktabs(t, len(rows)); set_cell_margins(t, top=50, bottom=50, left=90, right=90)
        numeric = [ci > 0 and all(looks_numeric(str(r[ci])) for r in rows) for ci in range(len(header))]
        total = sum(widths)
        widths = [w * min(1.0, TEXT_WIDTH_IN / total) for w in widths]
        for ri, row in enumerate([header] + rows):
            tr = t.rows[ri]; no_split(tr)
            if ri == 0: repeat_header(tr)
            for ci, val in enumerate(row):
                cell = tr.cells[ci]; cell.width = Inches(widths[ci])
                cp = cell.paragraphs[0]; cp.style = self.doc.styles["Table Text"]
                self.add_inline(cp, str(val), bold=(ri == 0))
                if numeric[ci]:
                    cp.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
                set_cell_valign(cell)
        for ci, w in enumerate(widths):
            for cell in t.columns[ci].cells:
                cell.width = Inches(w)
        set_grid_widths(t, widths)
        # small fixed gap after the table (a 2 pt paragraph), instead of a full blank line
        sp = self.doc.add_paragraph(style="Single"); sp.paragraph_format.space_after = Pt(10)
        sp.paragraph_format.line_spacing = Pt(2)

    def equation(self, key, latex):
        lab = self.num.labels[key]
        png = self.eq_dir / f"eq_{lab.replace('.', '_')}.png"
        w, h = render_equation(latex, png)
        p = self.doc.add_paragraph(style="Single")
        p.paragraph_format.space_before = Pt(8); p.paragraph_format.space_after = Pt(10)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(TEXT_WIDTH_IN / 2), WD_TAB_ALIGNMENT.CENTER)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(TAB_RIGHT_IN), WD_TAB_ALIGNMENT.RIGHT)
        p.add_run("\t")
        p.add_run().add_picture(str(png), width=Inches(min(w, 5.0)))
        p.add_run(f"\t({lab})")

    def bullets(self, items, numbered=False):
        for i, it in enumerate(items, 1):
            if numbered:
                p = self.doc.add_paragraph(style="Normal")
                p.paragraph_format.left_indent = Inches(0.4); p.paragraph_format.first_line_indent = Inches(-0.3)
                p.paragraph_format.space_after = Pt(4)
                p.add_run(f"{i}.\t")
                p.paragraph_format.tab_stops.add_tab_stop(Inches(0.4))
                self.add_inline(p, it)
            else:
                p = self.doc.add_paragraph(style="List Bullet")
                self.add_inline(p, it)
        self.doc.paragraphs[-1].paragraph_format.space_after = Pt(10)

    # ---------------- contents-type pages (Word refreshes them; the static text shows elsewhere)
    def contents_page(self, title, entries, instr, kind):
        h = self.doc.add_paragraph(style="Single")
        h.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER; h.paragraph_format.space_after = Pt(18)
        h.paragraph_format.page_break_before = True
        r = h.add_run(title); r.bold = True; r.font.size = Pt(14)
        first = True
        paras = []
        for e in entries:
            level, text, key = e
            p = self.doc.add_paragraph(style="TOC Entry")
            p.paragraph_format.tab_stops.add_tab_stop(Inches(TAB_RIGHT_IN), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
            if first and kind == "toc":
                field_begin(p, instr); first = False
            indent = {0: 0.0, 1: 0.0, 2: 0.35, 3: 0.75}[level]
            hang = {0: 0.95, 1: 0.0, 2: 0.45, 3: 0.55}[level]
            p.paragraph_format.left_indent = Inches(indent + hang)
            p.paragraph_format.first_line_indent = Inches(-hang)
            if kind == "toc" and level == 1:
                p.paragraph_format.space_before = Pt(8)
            run_text = self.resolve(text)
            clean = re.sub(r"[_^]\{(.+?)\}", r"\1", re.sub(r"\*\*|\*|`", "", run_text))
            if kind == "toc" and level == 1:
                clean = clean.replace("  ", ": ", 1)          # "Chapter 3: Study Area and Data"
            else:
                clean = clean.replace("  ", "\t", 1)         # number, tab, title (aligned column)
            p.paragraph_format.tab_stops.add_tab_stop(Inches(indent + hang))
            run = p.add_run(clean); run.bold = (kind == "toc" and level in (0, 1))
            if kind != "toc":
                run.font.size = Pt(11)
            p.add_run(f"\t{self.page_of(key)}")
            paras.append(p)
        if paras and kind == "toc":
            field_end(paras[-1])

    # ---------------- sections and footers
    def footer_page_numbers(self, section, show=True):
        section.footer.is_linked_to_previous = False
        fp = section.footer.paragraphs[0]
        for r in list(fp.runs):
            r._r.getparent().remove(r._r)
        if show:
            fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_field(fp, "PAGE", "1")
            for r in fp.runs:
                r.font.name = FONT; r.font.size = Pt(12)

    def new_section(self):
        s = self.doc.add_section(WD_SECTION.NEW_PAGE)
        s.left_margin, s.right_margin = Inches(1.5), Inches(1.0)
        s.top_margin, s.bottom_margin = Inches(1.0), Inches(1.0)
        return s

    # ---------------- main entry
    def _watermark(self, header, image, width_in, section=None):
        """Anchor a picture behind the text, centred on the page, in this header (Word's watermark mechanism)."""
        from docx.oxml import parse_xml
        from docx.oxml.ns import nsdecls
        par = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
        run = par.add_run()
        run.add_picture(str(image), width=Inches(width_in))
        inline = run._r.xpath(".//wp:inline")[0]
        ext = inline.find(qn("wp:extent"))
        doc_pr = inline.find(qn("wp:docPr"))
        # explicit offsets (not wp:align) so every viewer, not only Word, centres it on the page
        sec = section or self.doc.sections[-1]
        off_x = (int(sec.page_width) - int(ext.get("cx"))) // 2
        off_y = (int(sec.page_height) - int(ext.get("cy"))) // 2
        anchor = parse_xml(
            f'<wp:anchor {nsdecls("wp")} distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="0" '
            f'behindDoc="1" locked="1" layoutInCell="1" allowOverlap="1">'
            f'<wp:simplePos x="0" y="0"/>'
            f'<wp:positionH relativeFrom="page"><wp:posOffset>{off_x}</wp:posOffset></wp:positionH>'
            f'<wp:positionV relativeFrom="page"><wp:posOffset>{off_y}</wp:posOffset></wp:positionV>'
            f'<wp:extent cx="{ext.get("cx")}" cy="{ext.get("cy")}"/>'
            f'<wp:effectExtent l="0" t="0" r="0" b="0"/><wp:wrapNone/>'
            f'<wp:docPr id="{doc_pr.get("id")}" name="Watermark"/><wp:cNvGraphicFramePr/></wp:anchor>')
        anchor.append(inline.find(qn("a:graphic")))
        inline.getparent().replace(inline, anchor)

    def build(self, front_matter_fn, out_path, watermark=None, watermark_width_in=4.2):
        cover_section = self.doc.sections[0]
        front_matter_fn(self, stage="cover")
        fsec = self.new_section()
        front_matter_fn(self, stage="front")
        # the toc/lof/lot entries are known only after the body is laid out, so the body goes first into a
        # scratch pass: walk blocks to collect entries, then emit contents pages, then the body itself.
        self._collect_entries()
        self.contents_page("CONTENTS", self._toc_with_front(), 'TOC \\o "1-3" \\h \\z \\u', "toc")
        self.contents_page("LIST OF FIGURES", [(0, f"{a}  {b}", k) for a, b, k in self.lof_entries], 'TOC \\h \\z \\t "Figure Caption,1"', "lof")
        self.contents_page("LIST OF TABLES", [(0, f"{a}  {b}", k) for a, b, k in self.lot_entries], 'TOC \\h \\z \\t "Table Caption,1"', "lot")
        front_matter_fn(self, stage="symbols")
        bsec = self.new_section()
        self.toc_entries, self.lof_entries, self.lot_entries = [], [], []
        self._render_body()
        # page numbering: cover none; front matter roman (title page counted but not printed); body arabic from 1.
        # Section objects must be fetched by position now: python-docx re-points earlier Section objects at the
        # document-level sectPr whenever a new section is added.
        cover_section, fsec, bsec = self.doc.sections[0], self.doc.sections[1], self.doc.sections[2]
        self.footer_page_numbers(cover_section, show=False)
        self.footer_page_numbers(fsec, show=True)
        fsec.different_first_page_header_footer = True
        fsec.first_page_footer.is_linked_to_previous = False
        set_page_numbering(fsec, "lowerRoman", 1)
        self.footer_page_numbers(bsec, show=True)
        set_page_numbering(bsec, "decimal", 1)
        bsec.different_first_page_header_footer = False
        if watermark is not None:
            # every page except the cover (own section, no header) and the title page (first page of the
            # front-matter section, which gets an empty first-page header)
            fsec.header.is_linked_to_previous = False
            fsec.first_page_header.is_linked_to_previous = False
            self._watermark(fsec.header, watermark, watermark_width_in, fsec)
            bsec.header.is_linked_to_previous = False
            self._watermark(bsec.header, watermark, watermark_width_in, bsec)
        self.doc.core_properties.title = "Thesis report"
        self.doc.save(out_path)

    def _toc_with_front(self):
        front = [(0, t, k) for t, k in self.front_toc]
        return front + [(lvl, txt, key) for lvl, txt, key in self.toc_entries_all]

    def _collect_entries(self):
        self.toc_entries_all, self.lof_all, self.lot_all = [], [], []
        appendix = False
        for b in self.blocks:
            k = b[0]
            if k == "appendix_mode":
                appendix = True
            elif k == "chapter":
                lab = self.num.labels[b[1]]
                self.toc_entries_all.append((1, (f"Appendix {lab}  " if appendix else f"Chapter {lab}  ") + b[2], b[1]))
            elif k == "h2":
                self.toc_entries_all.append((2, f"{self.num.labels[b[1]]}  {b[2]}", b[1]))
            elif k == "h3":
                self.toc_entries_all.append((3, f"{self.num.labels[b[1]]}  {b[2]}", b[1]))
            elif k == "fig":
                self.lof_entries.append((f"Figure {self.num.labels[b[1]]}", list_title(b[4]), b[1]))
            elif k == "table":
                self.lot_entries.append((f"Table {self.num.labels[b[1]]}", list_title(b[2]), b[1]))
            elif k == "references":
                self.toc_entries_all.append((1, "References", "references"))

    def _render_body(self):
        appendix = False
        for b in self.blocks:
            k = b[0]
            if k == "appendix_mode":
                appendix = True
            elif k == "chapter":
                lab = self.num.labels[b[1]]
                self.heading(1, (f"Appendix {lab}  " if appendix else f"Chapter {lab}  ") + b[2], b[1])
            elif k == "h2":
                self.heading(2, f"{self.num.labels[b[1]]}  {b[2]}", b[1])
            elif k == "h3":
                self.heading(3, f"{self.num.labels[b[1]]}  {b[2]}", b[1])
            elif k == "p":
                self.para(b[1])
            elif k == "bullets":
                self.bullets(b[1])
            elif k == "numbered":
                self.bullets(b[1], numbered=True)
            elif k == "eq":
                self.equation(b[1], b[2])
            elif k == "fig":
                self.figure(b[1], b[2], b[3], b[4])
            elif k == "table":
                self.table(b[1], b[2], b[3], b[4], b[5])
            elif k == "pagebreak":
                self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            elif k == "references":
                self.references()
            else:
                raise ValueError(k)

    def references(self):
        p = self.doc.add_paragraph(style="Heading 1")
        p.add_run("REFERENCES")
        self.toc_entries.append((1, "References", "references"))
        for key in self.num.cite_order:
            n = self.num.cite_num[key]
            rp = self.doc.add_paragraph(style="Reference")
            rp.paragraph_format.tab_stops.add_tab_stop(Inches(0.45))
            rp.add_run(f"[{n}]\t")
            self.add_inline(rp, self.refs[key])


def list_title(caption):
    """Entry text for the list of figures / tables: an explicit 'Short|Full' title, else a shortened caption."""
    return caption.split("|", 1)[0] if "|" in caption else short_caption(caption)


def short_caption(caption, limit=74):
    """One-line title for the list of figures / tables: the caption's first sentence, cut at a clause boundary
    (before ' (' or ', ') when it would not fit on one line."""
    s = re.split(r"(?<=[a-z0-9\)])\. ", caption, maxsplit=1)[0].rstrip(".")
    plain = re.sub(r"[_^]\{(.+?)\}", r"\1", s)
    if len(plain) <= limit:
        return s
    for sep in (" (", "; "):
        parts = s.split(sep)
        out = parts[0]
        for part in parts[1:]:
            if len(re.sub(r"[_^]\{(.+?)\}", r"\1", out + sep + part)) > limit:
                break
            out += sep + part
        if len(re.sub(r"[_^]\{(.+?)\}", r"\1", out)) <= limit and out != s:
            return out
    return s
