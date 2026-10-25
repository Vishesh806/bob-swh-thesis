"""Front matter, laid out after the guideline's sample pages 1, 2, 4, 5 and 6."""
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.shared import Pt, Inches

TITLE = ("MULTI-STEP SIGNIFICANT WAVE HEIGHT FORECASTING IN THE BAY OF BENGAL "
         "USING HYBRID MACHINE LEARNING AND DEEP LEARNING FRAMEWORKS")
STUDENT = "Vishesh Bhagwandas Agarwal"
ROLL = "22NA30045"
SUPERVISOR = "Professor Swapnadip De Chowdhury"
DEGREE_LINES = ["Dual Degree (Bachelor of Technology and Master of Technology)",
                "in Ocean Engineering and Naval Architecture"]
DEPT = "DEPARTMENT OF OCEAN ENGINEERING AND NAVAL ARCHITECTURE"
INST = "INDIAN INSTITUTE OF TECHNOLOGY KHARAGPUR"
LOGO = "/Users/visheshagarwal/Desktop/bob_swh_thesis/assets/iitkgp_logo.png"
MONTH_YEAR = "November 2026"
YEAR = "2026"

from results_text import ABSTRACT  # written from the Phase 5 results at build time

KEYWORDS = ("significant wave height; multi-step forecasting; Bay of Bengal; tropical cyclones; ERA5; "
            "gradient boosting; LightGBM; CatBoost; sequence-to-sequence LSTM; attention; peak-weighted loss")

SYMBOLS = [
    ("H_{s}", "Significant wave height", "m"), ("T_{m}", "Mean wave period", "s"), ("θ_{w}", "Mean wave direction (coming from)", "°"),
    ("u_{10}, v_{10}", "Eastward and northward wind components at 10 m", "m/s"), ("U_{10}", "Wind speed at 10 m", "m/s"),
    ("θ_{U}", "Wind direction (coming from)", "°"), ("p", "Mean sea-level pressure", "hPa"), ("T_{2m}", "Air temperature at 2 m", "°C"),
    ("SST", "Sea surface temperature", "°C"), ("ρ_{a}", "Air density", "kg/m³"), ("C_{d}", "Drag coefficient at 10 m", "–"),
    ("τ", "Surface wind stress", "N/m²"), ("L_{0}", "Deep-water wavelength, gT²/2π", "m"), ("g", "Gravitational acceleration (9.81)", "m/s²"),
    ("R_{d}", "Specific gas constant of dry air (287.05)", "J/(kg·K)"), ("h", "Forecast lead time (horizon)", "h"),
    ("L", "Look-back window length", "h"), ("t", "Forecast origin", "–"), ("q", "Extreme-sea-state threshold (95th percentile)", "m"),
    ("β, γ", "Extra weights of the peak-weighted loss", "–"), ("τ_{s}", "Width of the sigmoid in the peak-weighted loss", "m"),
    ("SS", "Skill score relative to persistence", "–"),
]
ABBREVIATIONS = [
    ("ARIMA", "Autoregressive integrated moving average"), ("BoB", "Bay of Bengal"), ("Bi-LSTM", "Bidirectional long short-term memory network"),
    ("CC", "Correlation coefficient"), ("CDS", "Copernicus Climate Data Store"), ("DM", "Diebold–Mariano test"),
    ("DMS", "Direct multi-step (forecasting strategy)"), ("ECMWF", "European Centre for Medium-Range Weather Forecasts"),
    ("ERA5", "Fifth-generation ECMWF atmospheric reanalysis"), ("GBDT", "Gradient-boosted decision trees"),
    ("IMD", "India Meteorological Department"), ("INCOIS", "Indian National Centre for Ocean Information Services"),
    ("LSTM", "Long short-term memory network"), ("MAE", "Mean absolute error"), ("MSE", "Mean squared error"),
    ("MSLP", "Mean sea-level pressure"), ("NWP", "Numerical weather prediction"), ("OMNI", "Ocean Moored buoy Network for the northern Indian Ocean"),
    ("RMSE", "Root-mean-square error"), ("Seq2Seq", "Sequence-to-sequence (encoder–decoder)"), ("SI", "Scatter index"),
    ("SST", "Sea surface temperature"), ("TPE", "Tree-structured Parzen estimator"), ("WAM", "Wave Model (third-generation spectral model)"),
]


SMALL_WORDS = {"a", "an", "and", "as", "at", "by", "for", "from", "in", "of", "on", "or", "the", "to", "using", "with"}


def title_case(text):
    """Title case with short function words in lower case: 'Forecasting in the Bay of Bengal Using ...'."""
    words = text.lower().split()
    out = []
    for i, w in enumerate(words):
        if i > 0 and w in SMALL_WORDS - {"using"}:
            out.append(w)
        else:
            out.append("-".join(part[:1].upper() + part[1:] for part in w.split("-")))
    return " ".join(out)


def _centered(b, text, size=12, bold=False, italic=False, after=0, before=0, spacing=1.0):
    p = b.doc.add_paragraph(style="Single")
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(after); p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.line_spacing = spacing
    if text:
        b.add_inline(p, text, size=size, bold=bold, italic=italic)
    return p


def _page_title(b, text, first=False):
    p = b.doc.add_paragraph(style="Single")
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18); p.paragraph_format.keep_with_next = True
    if not first:
        p.paragraph_format.page_break_before = True
    r = p.add_run(text); r.bold = True; r.font.size = Pt(14)


def _logo(b, width_in=1.0):
    """IIT Kharagpur logo at the top-left of the page."""
    p = b.doc.add_paragraph(style="Single")
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(18)
    p.add_run().add_picture(LOGO, width=Inches(width_in))
    return p


def front_matter(b, stage):
    if stage == "cover":
        _logo(b)
        _centered(b, "", before=60)
        _centered(b, TITLE, size=16, bold=True, spacing=1.5, after=60)
        _centered(b, STUDENT, size=14, bold=True, italic=True)
        return

    if stage == "front":
        b.front_toc = []
        # ---- title page (page i, number not printed)
        _logo(b, 0.85)
        _centered(b, TITLE, size=16, bold=True, spacing=1.35, after=20, before=0)
        _centered(b, "Report submitted to", bold=True, italic=True, after=4)
        _centered(b, "Indian Institute of Technology Kharagpur", bold=True, italic=True, after=4)
        _centered(b, "for the award of the degree", bold=True, italic=True, after=12)
        _centered(b, "of", bold=True, italic=True, after=12)
        for line in DEGREE_LINES:
            _centered(b, line, size=14, bold=True, after=2)
        _centered(b, "by", bold=True, italic=True, before=12, after=12)
        _centered(b, f"{STUDENT} ({ROLL})", size=14, bold=True, after=12)
        _centered(b, "under the guidance of", italic=True, after=4)
        _centered(b, SUPERVISOR, bold=True, after=36)
        _centered(b, "DEPARTMENT OF OCEAN ENGINEERING", size=14, bold=True, after=0)
        _centered(b, "AND NAVAL ARCHITECTURE", size=14, bold=True, after=2)
        _centered(b, INST, size=14, bold=True, after=12)
        _centered(b, MONTH_YEAR, size=14, bold=True, after=24)
        _centered(b, f"© {YEAR} {STUDENT}. All rights reserved.", size=12, after=0)

        # ---- declaration
        _page_title(b, "DECLARATION"); b.front_toc.append(("Declaration", "front:declaration"))
        p = b.para("I certify that", style="Single"); p.paragraph_format.space_after = Pt(8)
        items = ["The work contained in this report is original and has been done by me under the guidance of my supervisor.",
                 "The work has not been submitted to any other Institute for any degree or diploma.",
                 "I have followed the guidelines provided by the Institute in preparing the report.",
                 "I have conformed to the norms and guidelines given in the Ethical Code of Conduct of the Institute.",
                 "Whenever I have used materials (data, theoretical analysis, figures, and text) from other sources, I have given due "
                 "credit to them by citing them in the text of the report and giving their details in the references. The ERA5 "
                 "reanalysis data used in this work were produced by ECMWF and are distributed by the Copernicus Climate Change "
                 "Service under the Creative Commons Attribution 4.0 licence, which permits reuse with attribution."]
        for letter, it in zip("abcde", items):
            q = b.doc.add_paragraph(style="Single"); q.paragraph_format.left_indent = Inches(0.4)
            q.paragraph_format.first_line_indent = Inches(-0.3); q.paragraph_format.space_after = Pt(8)
            q.paragraph_format.tab_stops.add_tab_stop(Inches(0.4))
            q.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            b.add_inline(q, f"({letter})\t{it}")
        s = b.para("Signature of the Student", style="Single", align=WD_ALIGN_PARAGRAPH.RIGHT); s.paragraph_format.space_before = Pt(60)
        b.para(f"{STUDENT} ({ROLL})", style="Single", align=WD_ALIGN_PARAGRAPH.RIGHT)
        dp = b.para("Date: 10/11/2026\tPlace: Kharagpur", style="Single", align=WD_ALIGN_PARAGRAPH.LEFT)
        dp.paragraph_format.space_before = Pt(24)
        dp.paragraph_format.tab_stops.add_tab_stop(Inches(5.95), WD_TAB_ALIGNMENT.RIGHT)

        # ---- certificate
        _page_title(b, "CERTIFICATE"); b.front_toc.append(("Certificate by the Supervisor", "front:certificate"))
        b.para(f"This is to certify that the project report entitled “{title_case(TITLE)}” submitted by {STUDENT} ({ROLL}) to "
               f"Indian Institute of Technology Kharagpur, India, is a record of bona fide project work carried out under my "
               f"supervision and guidance, and is worthy of consideration for the award of the degree of "
               f"{DEGREE_LINES[0]} {DEGREE_LINES[1]} of the Institute.", style="Normal")
        s = b.para("______________________", style="Single", align=WD_ALIGN_PARAGRAPH.RIGHT); s.paragraph_format.space_before = Pt(72)
        b.para(SUPERVISOR, style="Single", align=WD_ALIGN_PARAGRAPH.RIGHT)
        b.para("Department of Ocean Engineering and Naval Architecture", style="Single", align=WD_ALIGN_PARAGRAPH.RIGHT)
        b.para("Indian Institute of Technology Kharagpur", style="Single", align=WD_ALIGN_PARAGRAPH.RIGHT)
        d = b.para("Date: 10/11/2026", style="Single", align=WD_ALIGN_PARAGRAPH.LEFT); d.paragraph_format.space_before = Pt(24)

        # ---- acknowledgements
        _page_title(b, "ACKNOWLEDGEMENTS"); b.front_toc.append(("Acknowledgements", "front:ack"))
        for t in [
            f"I am deeply grateful to my supervisor, {SUPERVISOR}, Department of Ocean Engineering and Naval Architecture, IIT Kharagpur, "
            "for the guidance, encouragement and critical questions that shaped this work.",
            "I thank the faculty and staff of the Department of Ocean Engineering and Naval Architecture for the courses and facilities "
            "that made this project possible, and the members of the evaluation panel for their time and feedback.",
            "This work would not have been possible without open data. I thank the European Centre for Medium-Range Weather Forecasts "
            "and the Copernicus Climate Change Service for the ERA5 reanalysis, INCOIS for the OMNI buoy programme, and the developers "
            "of the open-source Python libraries used throughout: pandas, NumPy, xarray, SciPy, statsmodels, scikit-learn, LightGBM, "
            "CatBoost, XGBoost, Optuna, TensorFlow/Keras and Matplotlib.",
            "Finally, I thank my family and friends for their constant support.",
        ]:
            b.para(t, style="Normal")
        s = b.para(STUDENT, style="Single", align=WD_ALIGN_PARAGRAPH.RIGHT); s.paragraph_format.space_before = Pt(36)
        b.para("IIT Kharagpur, " + MONTH_YEAR, style="Single", align=WD_ALIGN_PARAGRAPH.RIGHT)

        # ---- abstract (single-spaced, 200-300 words)
        _page_title(b, "ABSTRACT"); b.front_toc.append(("Abstract", "front:abstract"))
        b.para(ABSTRACT, style="Single").paragraph_format.space_after = Pt(14)
        kw = b.doc.add_paragraph(style="Single")
        r = kw.add_run("Keywords: "); r.bold = True
        b.add_inline(kw, KEYWORDS)
        b.front_toc += [("Contents", "front:contents"), ("List of Figures", "front:lof"), ("List of Tables", "front:lot"),
                        ("List of Symbols and Abbreviations", "front:symbols")]
        return

    if stage == "symbols":
        p = b.doc.add_paragraph(style="Single")
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(18); p.paragraph_format.keep_with_next = True
        p.add_run().add_break(WD_BREAK.PAGE)
        r = p.add_run("LIST OF SYMBOLS AND ABBREVIATIONS"); r.bold = True; r.font.size = Pt(14)
        _sym_table(b, ["Symbol", "Description", "Unit"], SYMBOLS, [1.1, 3.6, 1.3])
        p = b.doc.add_paragraph(style="Single"); p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(12)
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.keep_with_next = True
        p.add_run().add_break(WD_BREAK.PAGE)
        r = p.add_run("ABBREVIATIONS"); r.bold = True; r.font.size = Pt(14)
        _sym_table(b, ["Abbreviation", "Meaning"], ABBREVIATIONS, [1.4, 4.6])


def _sym_table(b, header, rows, widths):
    from builder import set_booktabs, set_cell_margins, set_grid_widths
    t = b.doc.add_table(rows=1 + len(rows), cols=len(header))
    set_grid_widths(t, widths)
    set_booktabs(t, len(rows)); set_cell_margins(t, top=45, bottom=45, left=90, right=90)
    for ri, row in enumerate([header] + list(rows)):
        for ci, val in enumerate(row):
            cell = t.rows[ri].cells[ci]; cell.width = Inches(widths[ci])
            cp = cell.paragraphs[0]; cp.style = b.doc.styles["Table Text"]
            b.add_inline(cp, val, size=11, bold=(ri == 0))
            pass
